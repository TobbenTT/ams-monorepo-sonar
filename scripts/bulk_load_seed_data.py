#!/usr/bin/env python3
"""
Bulk load seed data from Excel files into SQLite database.
Only inserts into EMPTY tables. Skips tables that already have data.
"""

import sqlite3
import os
import json
import sys
from datetime import datetime, date

import openpyxl

DB_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'ocp_maintenance.db')
SEED_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'seed_data')
MAX_ROWS = 5000

def get_table_info(conn, table_name):
    c = conn.cursor()
    c.execute(f'SELECT COUNT(*) FROM "{table_name}"')
    count = c.fetchone()[0]
    c.execute(f'PRAGMA table_info("{table_name}")')
    cols = [(r[1], r[2]) for r in c.fetchall()]
    return count, cols

def table_exists(conn, table_name):
    c = conn.cursor()
    c.execute("SELECT name FROM sqlite_master WHERE type='table' AND name=?", (table_name,))
    return c.fetchone() is not None

def read_excel(filename, max_rows=None):
    path = os.path.join(SEED_DIR, filename)
    if not os.path.exists(path):
        print(f"  FILE NOT FOUND: {path}")
        return None, None
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows()
    header_row = next(rows_iter)
    headers = [cell.value for cell in header_row]
    rows = []
    for i, row in enumerate(rows_iter):
        if max_rows and i >= max_rows:
            break
        values = [cell.value for cell in row]
        if all(v is None for v in values):
            continue
        rows.append(values)
    wb.close()
    return headers, rows

def normalize_col(name):
    if name is None:
        return ''
    return name.lower().strip().replace(' ', '_').replace('-', '_').replace('.', '_')

def map_columns(excel_headers, db_columns):
    db_col_names = [c[0] for c in db_columns]
    db_norm = {normalize_col(c): c for c in db_col_names}
    mapping = []
    for i, eh in enumerate(excel_headers):
        en = normalize_col(eh)
        if en in db_norm:
            mapping.append((i, db_norm[en]))
        else:
            for dn, dc in db_norm.items():
                if en and dn and len(en) > 2 and len(dn) > 2 and (en in dn or dn in en):
                    mapping.append((i, dc))
                    break
    return mapping

def convert_value(val, col_type='TEXT'):
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.isoformat()
    if isinstance(val, date):
        return val.isoformat()
    if isinstance(val, (list, dict)):
        return json.dumps(val)
    return val

def insert_batch(conn, table_name, col_names, rows, batch_size=500):
    if not rows:
        return 0
    placeholders = ', '.join(['?' for _ in col_names])
    cols_str = ', '.join([f'"{c}"' for c in col_names])
    sql = f'INSERT OR IGNORE INTO "{table_name}" ({cols_str}) VALUES ({placeholders})'
    c = conn.cursor()
    total = 0
    for i in range(0, len(rows), batch_size):
        batch = rows[i:i+batch_size]
        try:
            c.executemany(sql, batch)
            conn.commit()
            total += len(batch)
        except Exception as e:
            conn.rollback()
            for row in batch:
                try:
                    c.execute(sql, row)
                    conn.commit()
                    total += 1
                except:
                    conn.rollback()
    return total

def create_table_if_missing(conn, table_name, ddl):
    if not table_exists(conn, table_name):
        print(f"  Creating {table_name} table")
        conn.cursor().execute(ddl)
        conn.commit()

def load_with_auto_map(conn, excel_file, table_name, max_rows_override=None):
    if not table_exists(conn, table_name):
        print(f"  TABLE NOT FOUND: {table_name}")
        return
    count, db_cols = get_table_info(conn, table_name)
    if count > 0:
        print(f"  SKIP {table_name} ({count} rows already)")
        return
    limit = max_rows_override or MAX_ROWS
    headers, rows = read_excel(excel_file, max_rows=limit)
    if not rows:
        print(f"  NO DATA in {excel_file}")
        return
    db_col_names = [c[0] for c in db_cols]
    mapping = map_columns(headers, db_cols)
    if not mapping:
        print(f"  NO COLUMN MAPPING found for {table_name}")
        return
    mapped_db_cols = [db_col for _, db_col in mapping]
    print(f"  Mapped {len(mapping)}/{len(db_col_names)} DB cols: {mapped_db_cols}")
    insert_rows = []
    for row in rows:
        vals = [None] * len(db_col_names)
        for excel_idx, db_col in mapping:
            db_idx = db_col_names.index(db_col)
            vals[db_idx] = convert_value(row[excel_idx])
        insert_rows.append(tuple(vals))
    total = insert_batch(conn, table_name, db_col_names, insert_rows)
    print(f"  {table_name}: inserted {total}/{len(rows)} rows")


# ---- Specialized loaders ----

def load_criticality_assessments(conn):
    count, db_cols = get_table_info(conn, 'criticality_assessments')
    if count > 0:
        print(f"  SKIP criticality_assessments ({count} rows)")
        return
    headers, rows = read_excel('02_criticality_assessment.xlsx')
    if not rows: return
    db_col_names = [c[0] for c in db_cols]
    insert_rows = []
    for i, row in enumerate(rows):
        rd = dict(zip(headers, row))
        criteria = json.dumps({
            'sap_func_loc': str(rd.get('sap_func_loc', '') or ''),
            'equipment_name': str(rd.get('equipment_name', '') or ''),
            'failure_event': str(rd.get('failure_event', '') or ''),
            'failure_consequence': str(rd.get('failure_consequence', '') or ''),
            'redundancy': str(rd.get('redundancy', '') or ''),
            'downtime_days': rd.get('downtime_days'),
            'process_loss_usd': rd.get('process_loss_usd'),
            'consequence_safety': rd.get('consequence_safety'),
            'consequence_environment': rd.get('consequence_environment'),
            'consequence_cost': rd.get('consequence_cost'),
            'consequence_rsc': rd.get('consequence_rsc'),
            'risk_coefficient': rd.get('risk_coefficient'),
            'risk_valuation': rd.get('risk_valuation'),
            'is_sce': rd.get('is_sce'),
        })
        insert_rows.append((
            f"CA-{i+1:05d}",
            str(rd.get('sap_func_loc', '') or ''),
            '2025-01-15T00:00:00',
            'OCP Engineering',
            'RCM-Criticality',
            criteria,
            rd.get('probability'),
            rd.get('risk_valuation'),
            str(rd.get('criticality_level', 'MEDIUM') or 'MEDIUM'),
            None, None, 'approved'
        ))
    total = insert_batch(conn, 'criticality_assessments', db_col_names, insert_rows)
    print(f"  criticality_assessments: inserted {total}/{len(rows)} rows")

def load_work_packages(conn):
    count, db_cols = get_table_info(conn, 'work_packages')
    if count > 0:
        print(f"  SKIP work_packages ({count} rows)")
        return
    headers, rows = read_excel('05_work_packages.xlsx')
    if not rows: return
    db_col_names = [c[0] for c in db_cols]
    wp_map = {}
    for row in rows:
        rd = dict(zip(headers, row))
        wpid = str(rd.get('wp_id', '') or '')
        if wpid not in wp_map:
            wp_map[wpid] = {'data': rd, 'operations': []}
        wp_map[wpid]['operations'].append({
            'operation_no': rd.get('operation_no'),
            'operation_desc': str(rd.get('operation_desc', '') or ''),
            'work_center': str(rd.get('work_center', '') or ''),
            'duration_hours': rd.get('duration_hours'),
            'num_persons': rd.get('num_persons'),
        })
    insert_rows = []
    for wpid, info in wp_map.items():
        rd = info['data']
        freq = str(rd.get('frequency', '') or '')
        freq_val, freq_unit = None, None
        if freq:
            parts = freq.strip().split()
            if len(parts) >= 2:
                try: freq_val = int(parts[0])
                except: pass
                freq_unit = parts[-1]
            else:
                freq_unit = freq
        insert_rows.append((
            wpid, str(rd.get('wp_desc', '') or ''), str(rd.get('task_list_group', '') or ''),
            str(rd.get('sap_func_loc', '') or ''), freq_val, freq_unit,
            None, None, str(rd.get('eqart', '') or ''), None, None,
            json.dumps([]), json.dumps(info['operations']), None, None, 'active'
        ))
    total = insert_batch(conn, 'work_packages', db_col_names, insert_rows)
    print(f"  work_packages: inserted {total}/{len(wp_map)} unique WPs from {len(rows)} rows")

def load_work_orders(conn):
    count, db_cols = get_table_info(conn, 'work_orders')
    if count > 0:
        print(f"  SKIP work_orders ({count} rows already)")
        return
    headers, rows = read_excel('06_work_order_history.xlsx', max_rows=MAX_ROWS)
    if not rows: return
    db_col_names = [c[0] for c in db_cols]
    insert_rows = []
    for row in rows:
        rd = dict(zip(headers, row))
        insert_rows.append((
            str(rd.get('order_number', '') or ''),
            str(rd.get('order_type', '') or ''),
            str(rd.get('sap_func_loc', '') or ''),
            str(rd.get('sap_func_loc_short', '') or ''),
            str(rd.get('priority', '') or ''),
            str(rd.get('status', '') or ''),
            convert_value(rd.get('actual_start') or rd.get('plan_start')),
            rd.get('total_hours'),
            str(rd.get('description', '') or ''),
            json.dumps({'total_cost_usd': rd.get('total_cost_usd'), 'work_center': str(rd.get('work_center', '') or '')})
        ))
    total = insert_batch(conn, 'work_orders', db_col_names, insert_rows)
    print(f"  work_orders: inserted {total}/{len(rows)} rows")

def load_shutdown_events(conn):
    count, db_cols = get_table_info(conn, 'shutdown_events')
    if count > 0:
        print(f"  SKIP shutdown_events ({count} rows)")
        return
    headers, rows = read_excel('08_shutdown_calendar.xlsx')
    if not rows: return
    db_col_names = [c[0] for c in db_cols]
    insert_rows = []
    for row in rows:
        rd = dict(zip(headers, row))
        dur = rd.get('duration_days', 0) or 0
        try: ph = float(dur) * 24
        except: ph = 0
        st = str(rd.get('status', '') or '').lower()
        comp = 100.0 if st == 'completed' else (50.0 if st == 'in_progress' else 0.0)
        wo = rd.get('work_orders_count', 0) or 0
        insert_rows.append((
            str(rd.get('shutdown_id', '') or ''), 'OCP-JRF',
            str(rd.get('description', '') or ''), str(rd.get('status', '') or ''),
            convert_value(rd.get('start_date')), convert_value(rd.get('end_date')),
            convert_value(rd.get('start_date')) if st in ('completed','in_progress') else None,
            convert_value(rd.get('end_date')) if st == 'completed' else None,
            ph, ph if st == 'completed' else None,
            wo, wo if st == 'completed' else 0, comp, 0, None, datetime.now().isoformat()
        ))
    total = insert_batch(conn, 'shutdown_events', db_col_names, insert_rows)
    print(f"  shutdown_events: inserted {total}/{len(rows)} rows")

def load_planning_kpi_snapshots(conn):
    count, db_cols = get_table_info(conn, 'planning_kpi_snapshots')
    if count > 0:
        print(f"  SKIP planning_kpi_snapshots ({count} rows)")
        return
    headers, rows = read_excel('12_planning_kpi_input.xlsx')
    if not rows: return
    db_col_names = [c[0] for c in db_cols]
    pm = {}
    for row in rows:
        rd = dict(zip(headers, row))
        p = str(rd.get('period', '') or '')
        pm.setdefault(p, []).append(rd)
    insert_rows = []
    for i, (period, records) in enumerate(pm.items()):
        kpis = {}
        on_t = below_t = 0
        for r in records:
            key = f"{r.get('area','')}_{r.get('planning_group','')}"
            comp = r.get('compliance_pct', 0) or 0
            kpis[key] = {k: r.get(k) for k in ['total_wo_planned','total_wo_executed','compliance_pct','backlog_hours','wrench_time_pct','schedule_compliance_pct','emergency_wo_pct']}
            try:
                if float(comp) >= 85: on_t += 1
                else: below_t += 1
            except: below_t += 1
        tot = on_t + below_t
        health = round(on_t/tot*100, 1) if tot > 0 else 0
        ps = f"{period}-01" if period and len(period) >= 7 else period
        insert_rows.append((f"PKS-{i+1:04d}", 'OCP-JRF', ps, ps, json.dumps(kpis), health, on_t, below_t, datetime.now().isoformat()))
    total = insert_batch(conn, 'planning_kpi_snapshots', db_col_names, insert_rows)
    print(f"  planning_kpi_snapshots: inserted {total}/{len(pm)} periods")

def load_de_kpi_snapshots(conn):
    count, db_cols = get_table_info(conn, 'de_kpi_snapshots')
    if count > 0:
        print(f"  SKIP de_kpi_snapshots ({count} rows)")
        return
    headers, rows = read_excel('13_de_kpi_input.xlsx')
    if not rows: return
    db_col_names = [c[0] for c in db_cols]
    pm = {}
    for row in rows:
        rd = dict(zip(headers, row))
        p = str(rd.get('period', '') or '')
        pm.setdefault(p, []).append(rd)
    insert_rows = []
    for i, (period, records) in enumerate(pm.items()):
        kpis = {}
        avail_sum = cnt = 0
        for r in records:
            area = str(r.get('area', '') or '')
            kpis[area] = {k: r.get(k) for k in ['mtbf_hours','mttr_hours','availability_pct','reliability_pct','oee_pct','failure_count','pm_compliance_pct','cost_per_ton_usd']}
            try:
                avail_sum += float(r.get('availability_pct', 0) or 0)
                cnt += 1
            except: pass
        avg = round(avail_sum/cnt, 1) if cnt > 0 else 0
        mat = 'Advanced' if avg >= 95 else ('Intermediate' if avg >= 85 else 'Developing')
        ps = f"{period}-01" if period and len(period) >= 7 else period
        insert_rows.append((f"DKS-{i+1:04d}", 'OCP-JRF', ps, ps, json.dumps(kpis), avg, avg, mat, datetime.now().isoformat()))
    total = insert_batch(conn, 'de_kpi_snapshots', db_col_names, insert_rows)
    print(f"  de_kpi_snapshots: inserted {total}/{len(pm)} periods")

def load_kpi_metrics(conn):
    count, db_cols = get_table_info(conn, 'kpi_metrics')
    if count > 0:
        print(f"  SKIP kpi_metrics ({count} rows)")
        return
    headers, rows = read_excel('13_de_kpi_input.xlsx')
    if not rows: return
    db_col_names = [c[0] for c in db_cols]
    insert_rows = []
    for i, row in enumerate(rows):
        rd = dict(zip(headers, row))
        period = str(rd.get('period', '') or '')
        ps = f"{period}-01" if period and len(period) >= 7 else period
        mtbf_h = rd.get('mtbf_hours', 0) or 0
        try: mtbf_d = round(float(mtbf_h)/24, 2)
        except: mtbf_d = 0
        insert_rows.append((
            f"KM-{i+1:05d}", 'OCP-JRF', str(rd.get('area', '') or ''),
            ps, ps, datetime.now().isoformat(), mtbf_d, rd.get('mttr_hours'),
            rd.get('availability_pct'), rd.get('oee_pct'), None, None,
            rd.get('pm_compliance_pct'), rd.get('failure_count'), None, None, None
        ))
    total = insert_batch(conn, 'kpi_metrics', db_col_names, insert_rows)
    print(f"  kpi_metrics: inserted {total}/{len(rows)} rows")

def load_maintenance_tasks(conn):
    count, db_cols = get_table_info(conn, 'maintenance_tasks')
    if count > 0:
        print(f"  SKIP maintenance_tasks ({count} rows)")
        return
    headers, rows = read_excel('16_route_sheets.xlsx', max_rows=MAX_ROWS)
    if not rows: return
    db_col_names = [c[0] for c in db_cols]
    insert_rows = []
    for i, row in enumerate(rows):
        rd = dict(zip(headers, row))
        dur = rd.get('duration_min', 0) or 0
        try: dh = round(float(dur)/60, 2)
        except: dh = 0
        insert_rows.append((
            f"RT-{str(rd.get('route_id',''))}-{str(rd.get('operation_no',''))}",
            None, str(rd.get('operation_desc', '') or ''), None,
            str(rd.get('task_list_type', '') or ''), 0, None, None, None, None, None, dh,
            None, None, 'route_sheet', None, None,
            json.dumps({'work_center': str(rd.get('work_center', '') or ''), 'sap_func_loc': str(rd.get('sap_func_loc', '') or '')}),
            None, None, None, 0, None, 'active'
        ))
    total = insert_batch(conn, 'maintenance_tasks', db_col_names, insert_rows)
    print(f"  maintenance_tasks: inserted {total}/{len(rows)} rows")

def load_backlog_items(conn):
    count, db_cols = get_table_info(conn, 'backlog_items')
    if count > 0:
        print(f"  SKIP backlog_items ({count} rows already)")
        return
    headers, rows = read_excel('23_active_backlog.xlsx')
    if not rows: return
    db_col_names = [c[0] for c in db_cols]
    insert_rows = []
    for row in rows:
        rd = dict(zip(headers, row))
        ms = str(rd.get('materials_status', '') or '').lower()
        mr = 1 if ms in ('available','ready','complete') else 0
        cd = str(rd.get('constraint_detail', '') or '').lower()
        sr = 1 if 'shutdown' in cd else 0
        insert_rows.append((
            str(rd.get('order_number', '') or ''), str(rd.get('notification_number', '') or ''),
            str(rd.get('sap_func_loc', '') or ''), str(rd.get('sap_func_loc_short', '') or ''),
            str(rd.get('priority', '') or ''), str(rd.get('order_type', '') or ''),
            str(rd.get('status', '') or ''), str(rd.get('waiting_reason', '') or ''),
            rd.get('estimated_hours'), str(rd.get('work_center', '') or ''), mr, sr,
            rd.get('backlog_age_days'), convert_value(rd.get('creation_date'))
        ))
    total = insert_batch(conn, 'backlog_items', db_col_names, insert_rows)
    print(f"  backlog_items: inserted {total}/{len(rows)} rows")

def load_failure_predictions(conn):
    count, db_cols = get_table_info(conn, 'failure_predictions')
    if count > 0:
        print(f"  SKIP failure_predictions ({count} rows)")
        return
    headers, rows = read_excel('30_reliability_data.xlsx', max_rows=MAX_ROWS)
    if not rows: return
    db_col_names = [c[0] for c in db_cols]
    em = {}
    for row in rows:
        rd = dict(zip(headers, row))
        key = str(rd.get('sap_func_loc', '') or '')
        if key and (key not in em or str(rd.get('period', '')) > str(em[key].get('period', ''))):
            em[key] = rd
    insert_rows = []
    for i, (key, rd) in enumerate(em.items()):
        avail = rd.get('availability_pct', 0) or 0
        rel = rd.get('reliability_pct', 0) or 0
        mtbf = rd.get('mtbf_hours', 0) or 0
        try: risk = round(max(0, min(100, 100 - float(avail))), 1)
        except: risk = 50
        try: pw = round(float(mtbf)/24, 0) if float(mtbf) > 0 else None
        except: pw = None
        wb = json.dumps({'mtbf_hours': mtbf, 'mttr_hours': rd.get('mttr_hours'), 'operating_hours': rd.get('operating_hours'), 'downtime_hours': rd.get('downtime_hours'), 'failure_count': rd.get('failure_count')})
        rec = 'Increase PM frequency' if risk > 30 else ('Continue current strategy' if risk > 15 else 'Consider extending intervals')
        insert_rows.append((
            f"FP-{i+1:05d}", key, str(rd.get('sap_func_loc_short', '') or ''),
            datetime.now().isoformat(), wb, None, rel, pw, 0.75, risk,
            str(rd.get('criticality', '') or ''), rec, 'active'
        ))
    total = insert_batch(conn, 'failure_predictions', db_col_names, insert_rows)
    print(f"  failure_predictions: inserted {total}/{len(em)} equipment")

def load_health_scores(conn):
    count, db_cols = get_table_info(conn, 'health_scores')
    if count > 0:
        print(f"  SKIP health_scores ({count} rows)")
        return
    headers, rows = read_excel('30_reliability_data.xlsx', max_rows=MAX_ROWS)
    if not rows: return
    db_col_names = [c[0] for c in db_cols]
    em = {}
    for row in rows:
        rd = dict(zip(headers, row))
        key = str(rd.get('sap_func_loc', '') or '')
        if key and (key not in em or str(rd.get('period', '')) > str(em[key].get('period', ''))):
            em[key] = rd
    insert_rows = []
    for i, (key, rd) in enumerate(em.items()):
        avail = rd.get('availability_pct', 0) or 0
        rel = rd.get('reliability_pct', 0) or 0
        mc = rd.get('maint_cost_usd', 0) or 0
        try: comp = round((float(avail) + float(rel))/2, 1)
        except: comp = 50
        hc = 'Excellent' if comp >= 95 else ('Good' if comp >= 85 else ('Fair' if comp >= 70 else 'Poor'))
        dims = json.dumps({'availability': float(avail) if avail else 0, 'reliability': float(rel) if rel else 0, 'maintenance_cost_usd': float(mc) if mc else 0})
        recs = []
        try:
            if float(avail) < 90: recs.append('Improve availability')
            if float(rel) < 90: recs.append('Address recurring failures')
        except: pass
        insert_rows.append((
            f"HS-{i+1:05d}", key, 'OCP-JRF', str(rd.get('sap_func_loc_short', '') or ''),
            datetime.now().isoformat(), dims, comp, hc, 'stable', json.dumps(recs) if recs else None
        ))
    total = insert_batch(conn, 'health_scores', db_col_names, insert_rows)
    print(f"  health_scores: inserted {total}/{len(em)} equipment")


def main():
    print("=" * 70)
    print("BULK SEED DATA LOADER")
    print("=" * 70)
    conn = sqlite3.connect(DB_PATH)
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA synchronous=NORMAL")
    print(f"\nDatabase: {os.path.abspath(DB_PATH)}")
    print(f"Seed dir: {os.path.abspath(SEED_DIR)}")
    print(f"Max rows per large file: {MAX_ROWS}\n")

    # === EXISTING EMPTY TABLES ===
    print("--- Loading into EXISTING empty tables ---")

    print("\n[02] Criticality assessments...")
    load_criticality_assessments(conn)

    print("\n[04] Measurement points...")
    create_table_if_missing(conn, 'measuring_points', """CREATE TABLE measuring_points (
        point_id TEXT PRIMARY KEY, sap_func_loc TEXT, sap_func_loc_short TEXT,
        equipment_name TEXT, equnr TEXT, point_desc TEXT, meas_type TEXT, unit TEXT,
        lower_limit REAL, upper_limit REAL, target_value REAL, reading_frequency TEXT,
        is_counter INTEGER, counter_overflow REAL)""")
    load_with_auto_map(conn, '04_measurement_points.xlsx', 'measuring_points')

    print("\n[05] Work packages...")
    load_work_packages(conn)

    print("\n[06] Work orders...")
    load_work_orders(conn)

    print("\n[08] Shutdown events...")
    load_shutdown_events(conn)

    print("\n[12] Planning KPI snapshots...")
    load_planning_kpi_snapshots(conn)

    print("\n[13] DE KPI snapshots...")
    load_de_kpi_snapshots(conn)

    print("\n[13b] KPI metrics...")
    load_kpi_metrics(conn)

    print("\n[16] Maintenance tasks...")
    load_maintenance_tasks(conn)

    print("\n[23] Backlog items...")
    load_backlog_items(conn)

    print("\n[30] Failure predictions...")
    load_failure_predictions(conn)

    print("\n[30b] Health scores...")
    load_health_scores(conn)

    # === NEW TABLES ===
    print("\n--- Creating + loading NEW tables ---")

    print("\n[07] Spare parts (full)...")
    create_table_if_missing(conn, 'spare_parts_full', """CREATE TABLE spare_parts_full (
        material_code TEXT PRIMARY KEY, sap_material_number TEXT, description TEXT,
        manufacturer TEXT, part_number TEXT, ved_class TEXT, fsn_class TEXT, abc_class TEXT,
        quantity_on_hand REAL, min_stock REAL, max_stock REAL, reorder_point REAL,
        lead_time_days INTEGER, unit_cost_usd REAL, unit_of_measure TEXT,
        applicable_equipment_csv TEXT, warehouse_location TEXT)""")
    load_with_auto_map(conn, '07_spare_parts_inventory.xlsx', 'spare_parts_full')

    print("\n[10] Field inspections...")
    create_table_if_missing(conn, 'field_inspections', """CREATE TABLE field_inspections (
        inspection_id INTEGER PRIMARY KEY AUTOINCREMENT, capture_id TEXT, sap_func_loc TEXT,
        sap_func_loc_short TEXT, equipment_name TEXT, equnr TEXT, capture_date TEXT,
        inspector TEXT, inspection_type TEXT, condition_score REAL, observation TEXT,
        photo_ref TEXT, gps_lat REAL, gps_lon REAL, follow_up_required INTEGER)""")
    load_with_auto_map(conn, '10_field_capture.xlsx', 'field_inspections')

    print("\n[11] Work centers...")
    create_table_if_missing(conn, 'work_centers', """CREATE TABLE work_centers (
        work_center TEXT PRIMARY KEY, description TEXT, plant TEXT, work_center_cat TEXT,
        person_responsible TEXT, capacity_hours_day REAL, num_workers INTEGER,
        cost_center TEXT, business_area TEXT)""")
    load_with_auto_map(conn, '11_work_centers.xlsx', 'work_centers')

    print("\n[15] Catalog profiles...")
    create_table_if_missing(conn, 'catalog_profiles', """CREATE TABLE catalog_profiles (
        profile_id INTEGER PRIMARY KEY AUTOINCREMENT, profile_code TEXT, profile_desc TEXT,
        catalog_type TEXT, catalog_code TEXT, equipment_name TEXT)""")
    count, _ = get_table_info(conn, 'catalog_profiles')
    if count == 0:
        headers, rows = read_excel('15_catalog_profiles.xlsx')
        if rows:
            db_cols_cp = ['profile_id','profile_code','profile_desc','catalog_type','catalog_code','equipment_name']
            ir = []
            for row in rows:
                rd = dict(zip(headers, row))
                ir.append((None,
                    str(rd.get('Cod. perfil catalogo', rd.get('Cód. perfil catálogo', '')) or ''),
                    str(rd.get('Denominacion perfil catalogo', rd.get('Denominación perfil catálogo', '')) or ''),
                    str(rd.get('Tipo catalogo', rd.get('Tipo catálogo', '')) or ''),
                    str(rd.get('Codigo catalogo', rd.get('Código catálogo', '')) or ''),
                    str(rd.get('equipment_name', '') or '')
                ))
            total = insert_batch(conn, 'catalog_profiles', db_cols_cp, ir)
            print(f"  catalog_profiles: inserted {total}/{len(rows)} rows")

    print("\n[17] Maintenance plans...")
    create_table_if_missing(conn, 'maintenance_plans', """CREATE TABLE maintenance_plans (
        plan_number TEXT PRIMARY KEY, plan_desc TEXT, plan_type TEXT, strategy TEXT,
        sap_func_loc TEXT, sap_func_loc_short TEXT, equipment_name TEXT, equnr TEXT,
        task_list_id TEXT, cycle_length TEXT, cycle_unit TEXT, start_date TEXT,
        call_horizon_pct REAL, scheduling_indicator TEXT, planning_group TEXT,
        work_center TEXT, status TEXT)""")
    count_mp, db_cols_mp = get_table_info(conn, 'maintenance_plans')
    if count_mp == 0:
        headers, rows = read_excel('17_maintenance_plans.xlsx')
        if rows:
            db_cn = [c[0] for c in db_cols_mp]
            mapping = map_columns(headers, db_cols_mp)
            ir = []
            seen = set()
            for row in rows:
                vals = [None] * len(db_cn)
                for ei, dc in mapping:
                    di = db_cn.index(dc)
                    vals[di] = convert_value(row[ei])
                pn = vals[0]
                if pn and pn not in seen:
                    seen.add(pn)
                    ir.append(tuple(vals))
            total = insert_batch(conn, 'maintenance_plans', db_cn, ir)
            print(f"  maintenance_plans: inserted {total}/{len(ir)} unique plans from {len(rows)} rows")

    print("\n[20] Settlement rules...")
    create_table_if_missing(conn, 'settlement_rules', """CREATE TABLE settlement_rules (
        rule_id INTEGER PRIMARY KEY AUTOINCREMENT, sap_func_loc TEXT, sap_func_loc_short TEXT,
        equipment_name TEXT, equnr TEXT, cost_center TEXT, wbs_element TEXT, asset_number TEXT,
        profit_center TEXT, business_area TEXT, settlement_rule TEXT, budget_annual_usd REAL)""")
    load_with_auto_map(conn, '20_financial_assignments.xlsx', 'settlement_rules')

    print("\n[24] Notifications raw...")
    create_table_if_missing(conn, 'notifications_raw', """CREATE TABLE notifications_raw (
        notification_number TEXT PRIMARY KEY, notif_type TEXT, description TEXT,
        sap_func_loc TEXT, sap_func_loc_short TEXT, equipment_name TEXT, equnr TEXT,
        priority TEXT, status TEXT, creation_date TEXT, malfunction_start TEXT,
        malfunction_end TEXT, reported_by TEXT, catalog_profile TEXT, damage_code TEXT,
        cause_code TEXT, order_number TEXT)""")
    load_with_auto_map(conn, '24_notifications.xlsx', 'notifications_raw')

    print("\n[25] Measurement documents...")
    create_table_if_missing(conn, 'measurement_documents', """CREATE TABLE measurement_documents (
        doc_id INTEGER PRIMARY KEY AUTOINCREMENT, doc_number TEXT, sap_func_loc TEXT,
        sap_func_loc_short TEXT, equipment_name TEXT, equnr TEXT, point_id TEXT,
        meas_date TEXT, meas_value REAL, unit TEXT, measured_by TEXT,
        is_alarm INTEGER, alarm_type TEXT)""")
    load_with_auto_map(conn, '25_measurement_documents.xlsx', 'measurement_documents')

    print("\n[26] Time logs...")
    create_table_if_missing(conn, 'time_logs', """CREATE TABLE time_logs (
        log_id INTEGER PRIMARY KEY AUTOINCREMENT, confirmation_number TEXT, order_number TEXT,
        operation TEXT, sap_func_loc TEXT, sap_func_loc_short TEXT, equipment_name TEXT,
        work_center TEXT, employee_id TEXT, conf_date TEXT, actual_hours REAL,
        activity_type TEXT, final_confirmation INTEGER)""")
    load_with_auto_map(conn, '26_time_confirmations.xlsx', 'time_logs')

    print("\n[27] Material movements...")
    create_table_if_missing(conn, 'material_movements', """CREATE TABLE material_movements (
        movement_id TEXT PRIMARY KEY, movement_type TEXT, movement_type_desc TEXT,
        movement_category TEXT, material_code TEXT, sap_material_number TEXT, description TEXT,
        quantity REAL, unit TEXT, order_number TEXT, sap_func_loc TEXT, sap_func_loc_short TEXT,
        equipment_name TEXT, posting_date TEXT, cost_center TEXT, storage_location_from TEXT,
        storage_location_to TEXT, total_value_usd REAL, document_number TEXT,
        reversal_indicator TEXT)""")
    load_with_auto_map(conn, '27_material_movements.xlsx', 'material_movements')

    print("\n[28] BOM items...")
    create_table_if_missing(conn, 'bom_items', """CREATE TABLE bom_items (
        bom_id INTEGER PRIMARY KEY AUTOINCREMENT, parent_func_loc TEXT, parent_func_loc_short TEXT,
        parent_equipment_name TEXT, parent_equnr TEXT, component_name TEXT, component_type TEXT,
        component_func_loc TEXT, component_func_loc_short TEXT, component_equipment_name TEXT,
        level INTEGER)""")
    load_with_auto_map(conn, '28_equipment_bom.xlsx', 'bom_items')

    print("\n[29] Cost history...")
    create_table_if_missing(conn, 'cost_history', """CREATE TABLE cost_history (
        cost_id INTEGER PRIMARY KEY AUTOINCREMENT, period TEXT, sap_func_loc TEXT,
        sap_func_loc_short TEXT, equipment_name TEXT, equnr TEXT, cost_element TEXT,
        cost_element_desc TEXT, value_category TEXT, amount_usd REAL, currency TEXT,
        order_number TEXT, posting_date TEXT)""")
    load_with_auto_map(conn, '29_cost_history.xlsx', 'cost_history')

    print("\n[30c] Reliability data (full)...")
    create_table_if_missing(conn, 'reliability_data', """CREATE TABLE reliability_data (
        id INTEGER PRIMARY KEY AUTOINCREMENT, sap_func_loc TEXT, sap_func_loc_short TEXT,
        equipment_name TEXT, equnr TEXT, period TEXT, operating_hours REAL, downtime_hours REAL,
        failure_count INTEGER, mtbf_hours REAL, mttr_hours REAL, availability_pct REAL,
        reliability_pct REAL, maint_cost_usd REAL, criticality TEXT)""")
    load_with_auto_map(conn, '30_reliability_data.xlsx', 'reliability_data')

    # === FINAL SUMMARY ===
    print("\n" + "=" * 70)
    print("FINAL TABLE STATUS")
    print("=" * 70)
    c = conn.cursor()
    c.execute("SELECT name FROM sqlite_master WHERE type='table' ORDER BY name")
    tables = [r[0] for r in c.fetchall()]
    total_rows = 0
    empty_count = 0
    for t in tables:
        c.execute(f'SELECT COUNT(*) FROM "{t}"')
        cnt = c.fetchone()[0]
        total_rows += cnt
        if cnt == 0:
            empty_count += 1
            print(f"  {t}: {cnt} (empty)")
        else:
            print(f"  {t}: {cnt} rows")
    print(f"\nTotal tables: {len(tables)}")
    print(f"Tables with data: {len(tables) - empty_count}")
    print(f"Empty tables: {empty_count}")
    print(f"Total rows: {total_rows}")
    conn.close()
    print("\nDone!")

if __name__ == '__main__':
    main()
