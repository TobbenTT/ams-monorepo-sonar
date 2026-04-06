#!/usr/bin/env python3
"""Seed 11 new Excel files (31-41) into ocp_maintenance.db"""
import sqlite3
import os
from datetime import datetime, date, time

try:
    import openpyxl
except ImportError:
    import subprocess, sys
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'openpyxl'])
    import openpyxl

DB_PATH = '/root/ASSET-MANAGEMENT-SOFTWARE/ocp_maintenance.db'
SEED_DIR = '/root/ASSET-MANAGEMENT-SOFTWARE/seed_data'

conn = sqlite3.connect(DB_PATH)
cur = conn.cursor()

TABLE_DEFS = {
    'maintenance_schedule_3w': """CREATE TABLE IF NOT EXISTS maintenance_schedule_3w (
        schedule_week TEXT, scheduled_date TEXT, order_number TEXT, order_type TEXT,
        sap_func_loc TEXT, sap_func_loc_short TEXT, equipment_name TEXT, order_description TEXT,
        priority TEXT, operation_number TEXT, operation_description TEXT, work_center TEXT,
        specialty TEXT, planned_hours REAL, num_persons INTEGER, planned_start TEXT,
        planned_end TEXT, material_required TEXT, material_available TEXT, status TEXT,
        source TEXT, constraint_notes TEXT
    )""",
    'typical_operations': """CREATE TABLE IF NOT EXISTS typical_operations (
        operation_catalog_id TEXT PRIMARY KEY, operation_category TEXT, operation_code TEXT,
        operation_description TEXT, description_detail TEXT, eqart_applicable TEXT,
        specialty_required TEXT, typical_duration_hours REAL, num_persons INTEGER,
        requires_permit TEXT, permit_type TEXT, safety_level TEXT, mandatory TEXT,
        sequence_order INTEGER
    )""",
    'shutdown_details': """CREATE TABLE IF NOT EXISTS shutdown_details (
        shutdown_id TEXT, activity_id TEXT, activity_name TEXT, activity_type TEXT,
        parent_activity_id TEXT, sap_func_loc TEXT, sap_func_loc_short TEXT,
        equipment_name TEXT, order_number TEXT, predecessor_id TEXT, dependency_type TEXT,
        lag_hours REAL, planned_start TEXT, planned_end TEXT, duration_hours REAL,
        work_center TEXT, specialty TEXT, num_persons INTEGER, is_critical_path TEXT,
        float_hours REAL, status TEXT, constraint_type TEXT, notes TEXT
    )""",
    'work_permits': """CREATE TABLE IF NOT EXISTS work_permits (
        permit_id TEXT PRIMARY KEY, permit_type TEXT, order_number TEXT, sap_func_loc TEXT,
        sap_func_loc_short TEXT, equipment_name TEXT, description TEXT, requested_by TEXT,
        approved_by TEXT, safety_supervisor TEXT, valid_from TEXT, valid_to TEXT, shift TEXT,
        loto_points INTEGER, loto_detail TEXT, gas_monitoring_required TEXT,
        rescue_plan_required TEXT, ppe_special TEXT, status TEXT, risk_level TEXT
    )""",
    'weekly_programs_seed': """CREATE TABLE IF NOT EXISTS weekly_programs_seed (
        program_week TEXT, program_date TEXT, planning_group TEXT, work_center TEXT,
        specialty TEXT, available_hh REAL, planned_hh_pm REAL, planned_hh_cm REAL,
        planned_hh_shutdown REAL, total_planned_hh REAL, utilization_pct REAL,
        overload_flag TEXT, num_orders_pm INTEGER, num_orders_cm INTEGER,
        order_number TEXT, order_type TEXT, order_description TEXT, notification_number TEXT,
        backlog_hh_remaining REAL, compliance_target_pct REAL, notes TEXT
    )""",
    'resource_availability': """CREATE TABLE IF NOT EXISTS resource_availability (
        date TEXT, shift TEXT, work_center TEXT, specialty TEXT, total_headcount INTEGER,
        available_headcount INTEGER, available_hh REAL, assigned_hh REAL, remaining_hh REAL,
        absence_reason TEXT, absent_count INTEGER
    )""",
    'material_reservations': """CREATE TABLE IF NOT EXISTS material_reservations (
        reservation_id TEXT PRIMARY KEY, order_number TEXT, operation_number TEXT,
        material_code TEXT, sap_material_number TEXT, description TEXT,
        quantity_required REAL, quantity_withdrawn REAL, quantity_available REAL,
        unit TEXT, storage_location TEXT, availability_status TEXT,
        expected_delivery_date TEXT, is_critical TEXT, ved_class TEXT
    )""",
    'kpi_snapshots_data': """CREATE TABLE IF NOT EXISTS kpi_snapshots_data (
        snapshot_date TEXT, period_type TEXT, area TEXT, planning_group TEXT,
        kpi_name TEXT, kpi_value REAL, kpi_target REAL, kpi_unit TEXT,
        trend TEXT, variance_pct REAL
    )""",
    'execution_checklists_data': """CREATE TABLE IF NOT EXISTS execution_checklists_data (
        checklist_id TEXT, eqart TEXT, maintenance_type TEXT, sap_func_loc TEXT,
        sap_func_loc_short TEXT, equipment_name TEXT, equnr TEXT, checklist_name TEXT,
        step_number INTEGER, step_description TEXT, step_type TEXT, expected_value TEXT,
        unit TEXT, is_mandatory TEXT, requires_photo TEXT, safety_critical TEXT
    )""",
    'daily_logs': """CREATE TABLE IF NOT EXISTS daily_logs (
        log_id TEXT PRIMARY KEY, log_date TEXT, shift TEXT, area TEXT, work_center TEXT,
        shift_supervisor TEXT, entry_type TEXT, description TEXT, related_order TEXT,
        sap_func_loc TEXT, sap_func_loc_short TEXT, equipment_name TEXT,
        priority_flag TEXT, follow_up_required TEXT, follow_up_assigned_to TEXT
    )""",
    'annual_budget_executive': """CREATE TABLE IF NOT EXISTS annual_budget_executive (
        kpi_category TEXT, kpi_name TEXT, kpi_description TEXT, unit TEXT,
        jan_budget REAL, feb_budget REAL, mar_budget REAL, apr_budget REAL,
        may_budget REAL, jun_budget REAL, jul_budget REAL, aug_budget REAL,
        sep_budget REAL, oct_budget REAL, nov_budget REAL, dec_budget REAL,
        annual_budget REAL, annual_actual REAL, variance_pct REAL
    )""",
    'annual_budget_opex': """CREATE TABLE IF NOT EXISTS annual_budget_opex (
        cost_center TEXT, cost_center_desc TEXT, business_area TEXT, cost_element TEXT,
        cost_element_desc TEXT, cost_sub_element TEXT,
        jan_budget_usd REAL, feb_budget_usd REAL, mar_budget_usd REAL, apr_budget_usd REAL,
        may_budget_usd REAL, jun_budget_usd REAL, jul_budget_usd REAL, aug_budget_usd REAL,
        sep_budget_usd REAL, oct_budget_usd REAL, nov_budget_usd REAL, dec_budget_usd REAL,
        annual_budget_usd REAL,
        jan_actual_usd REAL, feb_actual_usd REAL, mar_actual_usd REAL, apr_actual_usd REAL,
        may_actual_usd REAL, jun_actual_usd REAL, jul_actual_usd REAL, aug_actual_usd REAL,
        sep_actual_usd REAL, oct_actual_usd REAL, nov_actual_usd REAL, dec_actual_usd REAL,
        annual_actual_usd REAL, variance_usd REAL, variance_pct REAL
    )""",
    'annual_budget_maintenance': """CREATE TABLE IF NOT EXISTS annual_budget_maintenance (
        fleet_group TEXT, fleet_description TEXT, sap_func_loc_l2 TEXT, cost_center TEXT,
        planning_group TEXT, maint_cost_type TEXT, maint_cost_type_desc TEXT,
        jan_budget_usd REAL, feb_budget_usd REAL, mar_budget_usd REAL, apr_budget_usd REAL,
        may_budget_usd REAL, jun_budget_usd REAL, jul_budget_usd REAL, aug_budget_usd REAL,
        sep_budget_usd REAL, oct_budget_usd REAL, nov_budget_usd REAL, dec_budget_usd REAL,
        annual_budget_usd REAL, budget_source TEXT
    )""",
    'annual_budget_capex': """CREATE TABLE IF NOT EXISTS annual_budget_capex (
        capex_id TEXT, capex_category TEXT, capex_category_desc TEXT, project_name TEXT,
        wbs_element TEXT, cost_center TEXT, asset_class TEXT, sap_func_loc TEXT,
        jan_budget_usd REAL, feb_budget_usd REAL, mar_budget_usd REAL, apr_budget_usd REAL,
        may_budget_usd REAL, jun_budget_usd REAL, jul_budget_usd REAL, aug_budget_usd REAL,
        sep_budget_usd REAL, oct_budget_usd REAL, nov_budget_usd REAL, dec_budget_usd REAL,
        annual_budget_usd REAL, approval_status TEXT, priority TEXT
    )""",
    'annual_budget_kpi_targets': """CREATE TABLE IF NOT EXISTS annual_budget_kpi_targets (
        kpi_domain TEXT, kpi_name TEXT, kpi_description TEXT, area TEXT,
        planning_group TEXT, unit TEXT,
        jan_target REAL, feb_target REAL, mar_target REAL, apr_target REAL,
        may_target REAL, jun_target REAL, jul_target REAL, aug_target REAL,
        sep_target REAL, oct_target REAL, nov_target REAL, dec_target REAL,
        annual_target REAL, calculation_method TEXT, data_source_file TEXT
    )""",
    'annual_budget_equipment': """CREATE TABLE IF NOT EXISTS annual_budget_equipment (
        sap_func_loc TEXT, sap_func_loc_short TEXT, equipment_name TEXT, equnr TEXT,
        cost_center TEXT, planning_group TEXT, cost_element TEXT, cost_element_desc TEXT,
        jan_budget_usd REAL, feb_budget_usd REAL, mar_budget_usd REAL, apr_budget_usd REAL,
        may_budget_usd REAL, jun_budget_usd REAL, jul_budget_usd REAL, aug_budget_usd REAL,
        sep_budget_usd REAL, oct_budget_usd REAL, nov_budget_usd REAL, dec_budget_usd REAL,
        annual_budget_usd REAL
    )""",
    'annual_budget_production': """CREATE TABLE IF NOT EXISTS annual_budget_production (
        production_area TEXT, parameter TEXT, parameter_desc TEXT, unit TEXT,
        ene REAL, feb REAL, mar REAL, abr REAL, may REAL, jun REAL,
        jul REAL, ago REAL, sep REAL, oct REAL, nov REAL, dic REAL,
        annual_total REAL
    )""",
}

# Create all tables
for tname, ddl in TABLE_DEFS.items():
    cur.execute(ddl)
    print(f'[OK] Table ensured: {tname}')
conn.commit()

# File-to-table mapping: (filename, sheet_name, table_name)
FILE_MAP = [
    ('31_maintenance_schedule_3w.xlsx', 'Sheet1', 'maintenance_schedule_3w'),
    ('32_typical_operations.xlsx', 'Sheet1', 'typical_operations'),
    ('33_shutdown_detail.xlsx', 'Sheet1', 'shutdown_details'),
    ('34_permits_to_work.xlsx', 'Sheet1', 'work_permits'),
    ('35_weekly_program.xlsx', 'Sheet1', 'weekly_programs_seed'),
    ('36_resource_availability.xlsx', 'Sheet1', 'resource_availability'),
    ('37_material_reservations.xlsx', 'Material Reservations', 'material_reservations'),
    ('38_kpi_snapshots.xlsx', 'Sheet1', 'kpi_snapshots_data'),
    ('39_execution_checklists.xlsx', 'Sheet1', 'execution_checklists_data'),
    ('40_daily_log.xlsx', 'Sheet1', 'daily_logs'),
    ('41_annual_budget.xlsx', 'Executive_Dashboard', 'annual_budget_executive'),
    ('41_annual_budget.xlsx', 'OPEX_by_Area', 'annual_budget_opex'),
    ('41_annual_budget.xlsx', 'Maintenance_Budget', 'annual_budget_maintenance'),
    ('41_annual_budget.xlsx', 'CAPEX_Plan', 'annual_budget_capex'),
    ('41_annual_budget.xlsx', 'KPI_Targets', 'annual_budget_kpi_targets'),
    ('41_annual_budget.xlsx', 'Budget_Equipment_Detail', 'annual_budget_equipment'),
    ('41_annual_budget.xlsx', 'Production_Plan', 'annual_budget_production'),
]

def sanitize(val):
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.isoformat()
    if isinstance(val, date):
        return val.isoformat()
    if isinstance(val, time):
        return val.isoformat()
    return val

# Tables with PRIMARY KEY -> use INSERT OR REPLACE
PK_TABLES = {'typical_operations', 'work_permits', 'material_reservations', 'daily_logs'}

totals = {}
wb_cache = {}

for fname, sheet_name, table_name in FILE_MAP:
    fpath = os.path.join(SEED_DIR, fname)
    if not os.path.exists(fpath):
        print(f'[SKIP] {fname} not found')
        continue

    if fname not in wb_cache:
        wb_cache[fname] = openpyxl.load_workbook(fpath, read_only=True, data_only=True)
    wb = wb_cache[fname]
    ws = wb[sheet_name]

    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        print(f'[SKIP] {fname}/{sheet_name}: no data rows')
        continue

    headers = [str(h).strip() if h else f'col_{i}' for i, h in enumerate(rows[0])]
    header_map = {
        'Ene': 'ene', 'Feb': 'feb', 'Mar': 'mar', 'Abr': 'abr',
        'May': 'may', 'Jun': 'jun', 'Jul': 'jul', 'Ago': 'ago',
        'Sep': 'sep', 'Oct': 'oct', 'Nov': 'nov', 'Dic': 'dic'
    }
    headers = [header_map.get(h, h) for h in headers]

    ncols = len(headers)
    placeholders = ','.join(['?'] * ncols)
    col_names = ','.join(headers)

    if table_name in PK_TABLES:
        sql = f'INSERT OR REPLACE INTO {table_name} ({col_names}) VALUES ({placeholders})'
    else:
        cur.execute(f'DELETE FROM {table_name}')
        sql = f'INSERT INTO {table_name} ({col_names}) VALUES ({placeholders})'

    count = 0
    for row in rows[1:]:
        vals = [sanitize(v) for v in row[:ncols]]
        while len(vals) < ncols:
            vals.append(None)
        if all(v is None for v in vals):
            continue
        cur.execute(sql, vals)
        count += 1

    conn.commit()
    totals[table_name] = totals.get(table_name, 0) + count
    print(f'[LOADED] {fname} / {sheet_name} -> {table_name}: {count} rows')

for wb in wb_cache.values():
    wb.close()

conn.close()

print()
print('=' * 60)
print('SUMMARY - Total rows inserted per table:')
print('=' * 60)
for t, c in sorted(totals.items()):
    print(f'  {t:40s} {c:>6d} rows')
print(f'  {"TOTAL":40s} {sum(totals.values()):>6d} rows')
